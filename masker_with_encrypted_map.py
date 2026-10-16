"""masker_with_encrypted_map.py
批量脱敏并把 mapping 加密保存（AES-GCM，密码派生 PBKDF2）。
用法：
  python masker_with_encrypted_map.py mask <in_folder> <out_folder>
    - 脚本会提示输入加密口令（password），并把 mapping 写为 mapping_YYYYmmdd_HHMMSS.csv.enc

  python masker_with_encrypted_map.py unmask <in_folder> <out_folder> <mapping_enc_path>
    - 脚本会提示输入解密口令，并使用解密后的 mapping 执行还原。

依赖：
  pip install python-docx openpyxl python-pptx cryptography

注意：mapping 中包含明文原始值（在内存中解密），请确保口令安全并限制文件访问。脚本不会把明文 mapping 写到磁盘（除非你把解密后手动保存）。
"""

import os
import re
import sys
import csv
import io
import argparse
import getpass
from datetime import datetime
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.backends import default_backend
import base64

# 正则：邮箱、手机号（中国 11 位）、身份证（15 或 18）
EMAIL_RE = re.compile(r'([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)')
PHONE_RE = re.compile(r'(\b1[3-9]\d{9}\b)')
ID_RE = re.compile(r'\b(\d{15}|\d{17}[0-9Xx])\b')

TEMPLATES = {
    'EMAIL': '[EMAIL_{:05d}]',
    'PHONE': '[PHONE_{:05d}]',
    'ID': '[ID_{:05d}]'
}

# Encryption helpers
SALT_SIZE = 16
NONCE_SIZE = 12
KDF_ITERATIONS = 390000

def derive_key(password: str, salt: bytes) -> bytes:
    password_bytes = password.encode('utf-8')
    kdf = PBKDF2HMAC(
        algorithm=hashes.SHA256(),
        length=32,
        salt=salt,
        iterations=KDF_ITERATIONS,
        backend=default_backend()
    )
    key = kdf.derive(password_bytes)
    return key

def encrypt_bytes(password: str, data: bytes) -> bytes:
    salt = os.urandom(SALT_SIZE)
    key = derive_key(password, salt)
    aesgcm = AESGCM(key)
    nonce = os.urandom(NONCE_SIZE)
    ct = aesgcm.encrypt(nonce, data, None)
    # store: b'ENCMAP' + salt + nonce + ct
    return b'ENCMAP' + salt + nonce + ct

def decrypt_bytes(password: str, enc: bytes) -> bytes:
    if not enc.startswith(b'ENCMAP'):
        raise ValueError('Unknown encrypted mapping format')
    body = enc[len(b'ENCMAP'):]
    salt = body[:SALT_SIZE]
    nonce = body[SALT_SIZE:SALT_SIZE+NONCE_SIZE]
    ct = body[SALT_SIZE+NONCE_SIZE:]
    key = derive_key(password, salt)
    aesgcm = AESGCM(key)
    pt = aesgcm.decrypt(nonce, ct, None)
    return pt

class Mapper:
    def __init__(self):
        self.maps = {'EMAIL': {}, 'PHONE': {}, 'ID': {}}
        self.cnts = {'EMAIL': 0, 'PHONE': 0, 'ID': 0}
        self.first_seen = {}

    def get_mask(self, typ, original, source_file=None):
        d = self.maps[typ]
        if original in d:
            return d[original]
        self.cnts[typ] += 1
        token = TEMPLATES[typ].format(self.cnts[typ])
        d[original] = token
        if token not in self.first_seen:
            self.first_seen[token] = source_file or ''
        return token

    def save_csv_bytes(self):
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(['masked', 'original', 'type', 'first_seen_in_file'])
        for typ in ['EMAIL', 'PHONE', 'ID']:
            for orig, masked in self.maps[typ].items():
                w.writerow([masked, orig, typ, self.first_seen.get(masked, '')])
        return buf.getvalue().encode('utf-8')

    @classmethod
    def load_from_csv_bytes(cls, b: bytes):
        s = b.decode('utf-8')
        buf = io.StringIO(s)
        r = csv.DictReader(buf)
        m = cls()
        for row in r:
            masked = row['masked']
            original = row['original']
            typ = row['type']
            m.maps.setdefault(typ, {})[original] = masked
            mo = re.search(r'_(\d+)\]', masked)
            if mo:
                n = int(mo.group(1))
                m.cnts[typ] = max(m.cnts.get(typ, 0), n)
            m.first_seen[masked] = row.get('first_seen_in_file', '')
        return m


def mask_text(text, mapper, source_file=None):
    if not text:
        return text
    def e_repl(m):
        orig = m.group(1)
        tok = mapper.get_mask('EMAIL', orig, source_file)
        return tok
    text = EMAIL_RE.sub(e_repl, text)
    def p_repl(m):
        orig = m.group(1)
        tok = mapper.get_mask('PHONE', orig, source_file)
        return tok
    text = PHONE_RE.sub(p_repl, text)
    def id_repl(m):
        orig = m.group(1)
        tok = mapper.get_mask('ID', orig, source_file)
        return tok
    text = ID_RE.sub(id_repl, text)
    return text


def unmask_text(text, mapper):
    if not text:
        return text
    rev = {}
    for typ in mapper.maps:
        for orig, masked in mapper.maps[typ].items():
            rev[masked] = orig
    if not rev:
        return text
    tokens = sorted(rev.keys(), key=len, reverse=True)
    for t in tokens:
        if t in text:
            text = text.replace(t, rev[t])
    return text


def process_docx(in_path, out_path, mapper=None, mode='mask'):
    doc = Document(in_path)
    for p in doc.paragraphs:
        for run in p.runs:
            if mode == 'mask':
                run.text = mask_text(run.text, mapper, in_path)
            else:
                run.text = unmask_text(run.text, mapper)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if mode == 'mask':
                    cell.text = mask_text(cell.text, mapper, in_path)
                else:
                    cell.text = unmask_text(cell.text, mapper)
    doc.save(out_path)


def process_xlsx(in_path, out_path, mapper=None, mode='mask'):
    wb = load_workbook(in_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if mode == 'mask':
                        cell.value = mask_text(cell.value, mapper, in_path)
                    else:
                        cell.value = unmask_text(cell.value, mapper)
    wb.save(out_path)


def process_pptx(in_path, out_path, mapper=None, mode='mask'):
    prs = Presentation(in_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if mode == 'mask':
                        run.text = mask_text(run.text, mapper, in_path)
                    else:
                        run.text = unmask_text(run.text, mapper)
    prs.save(out_path)


def ensure_dir(path):
    if not os.path.exists(path):
        os.makedirs(path)


def walk_and_process(in_folder, out_folder, mapper, mode='mask'):
    ensure_dir(out_folder)
    exts = {'.docx': process_docx, '.xlsx': process_xlsx, '.pptx': process_pptx}
    for root, _, files in os.walk(in_folder):
        rel = os.path.relpath(root, in_folder)
        out_sub = os.path.join(out_folder, rel) if rel != '.' else out_folder
        ensure_dir(out_sub)
        for f in files:
            name, ext = os.path.splitext(f)
            ext = ext.lower()
            if ext in exts:
                in_path = os.path.join(root, f)
                out_path = os.path.join(out_sub, f"{name}{('_masked' if mode=='mask' else '_unmasked')}{ext}")
                try:
                    print(f"Processing {in_path} -> {out_path}")
                    exts[ext](in_path, out_path, mapper=mapper, mode=mode)
                except Exception as e:
                    print(f"Failed {in_path}: {e}")


def cmd_mask(in_folder, out_folder):
    ensure_dir(out_folder)
    mapper = Mapper()
    walk_and_process(in_folder, out_folder, mapper, mode='mask')
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    csv_name = f'mapping_{ts}.csv'
    csv_bytes = mapper.save_csv_bytes()
    # ask password
    pw = getpass.getpass('请输入用于加密 mapping 的口令 (不会回显): ')
    if not pw:
        print('未输入口令，取消')
        return
    enc = encrypt_bytes(pw, csv_bytes)
    enc_name = csv_name + '.enc'
    enc_path = os.path.join(out_folder, enc_name)
    with open(enc_path, 'wb') as f:
        f.write(enc)
    print(f'Encrypted mapping saved to: {enc_path}')


def cmd_unmask(in_folder, out_folder, mapping_enc_path):
    ensure_dir(out_folder)
    if not os.path.exists(mapping_enc_path):
        print('找不到 mapping 文件:', mapping_enc_path)
        return
    pw = getpass.getpass('请输入用于解密 mapping 的口令 (不会回显): ')
    if not pw:
        print('未输入口令，取消')
        return
    with open(mapping_enc_path, 'rb') as f:
        enc = f.read()
    try:
        csv_bytes = decrypt_bytes(pw, enc)
    except Exception as e:
        print('解密失败:', e)
        return
    mapper = Mapper.load_from_csv_bytes(csv_bytes)
    walk_and_process(in_folder, out_folder, mapper, mode='unmask')
    print('Unmask completed')


def main():
    parser = argparse.ArgumentParser(description='Batch mask/unmask with encrypted mapping')
    subparsers = parser.add_subparsers(dest='cmd')

    p_mask = subparsers.add_parser('mask')
    p_mask.add_argument('in_folder')
    p_mask.add_argument('out_folder')

    p_unmask = subparsers.add_parser('unmask')
    p_unmask.add_argument('in_folder')
    p_unmask.add_argument('out_folder')
    p_unmask.add_argument('mapping_enc_path')

    args = parser.parse_args()
    if args.cmd == 'mask':
        cmd_mask(args.in_folder, args.out_folder)
    elif args.cmd == 'unmask':
        cmd_unmask(args.in_folder, args.out_folder, args.mapping_enc_path)
    else:
        parser.print_help()

if __name__ == '__main__':
    main()
